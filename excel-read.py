from openpyxl import load_workbook

workbook = load_workbook("AUTO-AERO.xlsx")
sheet = workbook.active

for row in sheet.iter_rows(min_row=1, max_row=3, max_col=2, values_only=True):
    print(row)
